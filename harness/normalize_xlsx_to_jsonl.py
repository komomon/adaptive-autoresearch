#!/usr/bin/env python3
"""
Normalize an evaluation xlsx sheet into JSONL cases for Adaptive Autoresearch.

This script is intentionally small and reusable:
- it does not know any specific vulnerability type
- it maps worksheet columns into the shared NormalizedCase contract
- it can be reused by authz, traditional sink, and business-logic datasets
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path
from typing import Any, Dict, Iterable, List, Tuple


REQUIRED_TARGETS = (
    "repo.url",
    "repo.branch",
    "entry.name",
    "expected.has_vulnerability",
)

VALID_TRANSPORTS = {
    "http",
    "rpc",
    "cli",
    "mq",
    "scheduler",
    "file",
    "internal-service",
    "other",
}

BOOL_TRUE = {"1", "true", "yes", "y", "是", "有", "存在"}
BOOL_FALSE = {"0", "false", "no", "n", "否", "无", "不存在"}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Convert xlsx cases into NormalizedCase JSONL."
    )
    parser.add_argument("--input", required=True, help="Path to the input xlsx file.")
    parser.add_argument("--output", required=True, help="Path to the output jsonl file.")
    parser.add_argument(
        "--sheet",
        default="cases",
        help="Worksheet name. Defaults to 'cases'.",
    )
    parser.add_argument(
        "--map",
        action="append",
        default=[],
        metavar="TARGET=SOURCE",
        help="Required field mapping. Example: repo.url=repo_url",
    )
    parser.add_argument(
        "--optional-map",
        action="append",
        default=[],
        metavar="TARGET=SOURCE",
        help="Optional field mapping. Example: expected.severity=severity",
    )
    parser.add_argument(
        "--split-multi-line",
        action="append",
        default=["expected.vuln_function_or_line"],
        metavar="TARGET",
        help="Fields that should split multi-line cells into string arrays.",
    )
    parser.add_argument(
        "--strict-transport",
        action="store_true",
        help="Fail if transport is outside the shared enum.",
    )
    parser.add_argument(
        "--default-transport",
        default="other",
        help="Fallback transport for sparse rows. Defaults to 'other'.",
    )
    parser.add_argument(
        "--default-language",
        default="unknown",
        help="Fallback language for sparse rows. Defaults to 'unknown'.",
    )
    return parser.parse_args()


def require_openpyxl():
    try:
        from openpyxl import load_workbook  # type: ignore
    except ImportError as exc:  # pragma: no cover
        raise SystemExit(
            "openpyxl is required to read .xlsx files. "
            "Install it in your environment before running this script."
        ) from exc
    return load_workbook


def parse_mapping_items(items: Iterable[str]) -> Dict[str, str]:
    mappings: Dict[str, str] = {}
    for item in items:
        if "=" not in item:
            raise SystemExit(f"Invalid mapping '{item}'. Expected TARGET=SOURCE.")
        target, source = item.split("=", 1)
        target = target.strip()
        source = source.strip()
        if not target or not source:
            raise SystemExit(f"Invalid mapping '{item}'. Target and source are required.")
        mappings[target] = source
    return mappings


def normalize_header(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def normalize_scalar(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, str):
        trimmed = value.strip()
        return trimmed if trimmed else None
    return value


def normalize_bool(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    text = str(value).strip().lower()
    if text in BOOL_TRUE:
        return True
    if text in BOOL_FALSE:
        return False
    raise ValueError(f"Cannot normalize boolean value: {value!r}")


def normalize_transport(value: Any, strict: bool) -> str:
    text = str(value).strip().lower()
    if not text:
        raise ValueError("transport cannot be empty")
    if text in VALID_TRANSPORTS:
        return text
    if strict:
        raise ValueError(f"Unsupported transport value: {value!r}")
    return "other"


def normalize_multiline(value: Any) -> Any:
    text = normalize_scalar(value)
    if text is None:
        return None
    assert isinstance(text, str)
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    if len(lines) <= 1:
        return text
    return lines


def set_nested(data: Dict[str, Any], dotted_key: str, value: Any) -> None:
    parts = dotted_key.split(".")
    current = data
    for part in parts[:-1]:
        if part not in current:
            current[part] = {}
        current = current[part]
    current[parts[-1]] = value


def get_nested(data: Dict[str, Any], dotted_key: str) -> Any:
    current: Any = data
    for part in dotted_key.split("."):
        if not isinstance(current, dict) or part not in current:
            return None
        current = current[part]
    return current


def validate_required_mappings(required_mappings: Dict[str, str]) -> None:
    missing = [key for key in REQUIRED_TARGETS if key not in required_mappings]
    if missing:
        joined = ", ".join(missing)
        raise SystemExit(f"Missing required mappings: {joined}")

def read_rows(
    input_path: Path,
    sheet_name: str,
) -> Tuple[List[str], List[Tuple[int, List[Any]]]]:
    load_workbook = require_openpyxl()
    workbook = load_workbook(filename=input_path, read_only=True, data_only=True)
    try:
        worksheet = workbook[sheet_name]
    except KeyError as exc:
        available = ", ".join(workbook.sheetnames)
        raise SystemExit(
            f"Worksheet '{sheet_name}' not found. Available sheets: {available}"
        ) from exc

    rows = list(worksheet.iter_rows(values_only=True))
    if not rows:
        raise SystemExit("The worksheet is empty.")

    headers = [normalize_header(cell) for cell in rows[0]]
    payload_rows = [(row_index + 2, list(row)) for row_index, row in enumerate(rows[1:])]
    return headers, payload_rows


def convert_row(
    row_number: int,
    row: List[Any],
    header_index: Dict[str, int],
    required_mappings: Dict[str, str],
    optional_mappings: Dict[str, str],
    split_multi_line_targets: List[str],
    strict_transport: bool,
    default_transport: str,
    default_language: str,
) -> Dict[str, Any]:
    payload: Dict[str, Any] = {}
    all_mappings = {**required_mappings, **optional_mappings}
    inferred_fields: List[str] = []
    missing_fields: List[str] = []

    for target, source in all_mappings.items():
        if source not in header_index:
            if target in required_mappings:
                raise ValueError(
                    f"Row {row_number}: source column '{source}' not found in worksheet."
                )
            continue

        raw_value = row[header_index[source]] if header_index[source] < len(row) else None
        value = normalize_scalar(raw_value)
        if value is None:
            continue

        if target == "expected.has_vulnerability":
            value = normalize_bool(value)
        elif target == "entry.transport":
            value = normalize_transport(value, strict_transport)
        elif target in split_multi_line_targets:
            value = normalize_multiline(value)

        set_nested(payload, target, value)

    if get_nested(payload, "case_id") is None:
        payload["case_id"] = f"row-{row_number:05d}"
        inferred_fields.append("case_id")

    if get_nested(payload, "entry.transport") is None:
        set_nested(payload, "entry.transport", normalize_transport(default_transport, strict=False))
        inferred_fields.append("entry.transport")
        missing_fields.append("entry.transport")

    if get_nested(payload, "entry.language") is None:
        set_nested(payload, "entry.language", default_language)
        inferred_fields.append("entry.language")
        missing_fields.append("entry.language")

    if get_nested(payload, "expected.vulnerability_type") is None:
        set_nested(payload, "expected.vulnerability_type", "unknown")
        inferred_fields.append("expected.vulnerability_type")
        missing_fields.append("expected.vulnerability_type")

    if get_nested(payload, "expected.vuln_function_or_line") is None:
        missing_fields.append("expected.vuln_function_or_line")

    payload["normalization"] = {
        "source_row": row_number,
        "inferred_fields": inferred_fields,
        "missing_fields": missing_fields,
        "needs_llm_review": bool(missing_fields),
    }

    return payload


def validate_required_fields(case: Dict[str, Any], row_number: int) -> None:
    missing = [target for target in REQUIRED_TARGETS if get_nested(case, target) is None]
    if missing:
        joined = ", ".join(missing)
        raise ValueError(f"Row {row_number}: missing required values for {joined}")


def main() -> int:
    args = parse_args()
    required_mappings = parse_mapping_items(args.map)
    optional_mappings = parse_mapping_items(args.optional_map)
    validate_required_mappings(required_mappings)

    input_path = Path(args.input).resolve()
    output_path = Path(args.output).resolve()
    output_path.parent.mkdir(parents=True, exist_ok=True)

    headers, payload_rows = read_rows(input_path, args.sheet)
    header_index = {header: idx for idx, header in enumerate(headers) if header}

    cases: List[Dict[str, Any]] = []
    errors: List[str] = []

    for row_number, row in payload_rows:
        if all(normalize_scalar(cell) is None for cell in row):
            continue
        try:
            case = convert_row(
                row_number=row_number,
                row=row,
                header_index=header_index,
                required_mappings=required_mappings,
                optional_mappings=optional_mappings,
                split_multi_line_targets=args.split_multi_line,
                strict_transport=args.strict_transport,
                default_transport=args.default_transport,
                default_language=args.default_language,
            )
            validate_required_fields(case, row_number)
            cases.append(case)
        except Exception as exc:  # pragma: no cover
            errors.append(str(exc))

    if errors:
        raise SystemExit("Normalization failed:\n- " + "\n- ".join(errors))

    with output_path.open("w", encoding="utf-8") as handle:
        for case in cases:
            handle.write(json.dumps(case, ensure_ascii=False) + "\n")

    print(f"Wrote {len(cases)} normalized cases to {output_path}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
